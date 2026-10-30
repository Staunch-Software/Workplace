"""15 ppm bilge alarm / oil content monitor (OCM) records for alert 23.

The monitor stores date/time, alarm status and operating status. Makers export these differently, so the import
is tolerant: it looks for a date/time column, a status column and (optionally) a ppm column by their header text,
and reduces every row to "was the vessel discharging overboard at this time". The status wording is matched by
keyword (see classify); rows whose status is neither clearly discharging nor clearly not are left out and counted.
Times are read day-first (dd/mm/yyyy) and are taken as ship local time, the clock the Oil Record Book uses.
Once a real export of each make is at hand the header and status lists below are the only things to adjust.
"""
import csv
import io
import re
from dataclasses import dataclass
from datetime import datetime, timedelta

_TIME_HEADERS = ("date/time", "date time", "datetime", "timestamp", "time stamp", "date and time", "date", "time")
_STATUS_HEADERS = ("operating status", "operation status", "status", "state", "mode", "operation", "event", "alarm status", "alarm")
_PPM_HEADERS = ("ppm", "oil content", "oil in water", "content", "oil")

_NOT_DISCHARGING = re.compile(r"recirc|stop|\boff\b|standby|stand-by|fault|closed|idle|flush|clean|calib|zero", re.I)
_DISCHARGING = re.compile(r"discharg|overboard|dump|running|\bon\b|\bopen", re.I)

_FORMATS = (
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M",
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S", "%d-%b-%Y %H:%M:%S", "%d-%b-%Y %H:%M", "%Y-%m-%dT%H:%M:%S",
)


def parse_ts(value) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    s = str(value or "").strip()
    if not s:
        return None
    for fmt in _FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def classify(status_text) -> bool | None:
    """True = discharging overboard, False = not, None = cannot tell from the wording."""
    t = str(status_text or "")
    if not t.strip():
        return None
    if _NOT_DISCHARGING.search(t):
        return False
    if _DISCHARGING.search(t):
        return True
    return None


def _pick(headers: list[str], wanted: tuple) -> str | None:
    low = {h: (h or "").strip().lower() for h in headers}
    for w in wanted:                                  # earlier entries of `wanted` win
        for h, l in low.items():
            if l == w:
                return h
    for w in wanted:
        for h, l in low.items():
            if w in l:
                return h
    return None


@dataclass
class ParsedOcm:
    records: list            # (recorded_at, discharging, ppm, status_text)
    rows_read: int
    rows_skipped: int
    headers: list
    time_col: str | None
    status_col: str | None


def parse_rows(rows: list[dict]) -> ParsedOcm:
    headers = list(rows[0].keys()) if rows else []
    time_col = _pick(headers, _TIME_HEADERS)
    status_col = _pick([h for h in headers if h != time_col], _STATUS_HEADERS)
    ppm_col = _pick([h for h in headers if h not in (time_col, status_col)], _PPM_HEADERS)
    records, skipped = [], 0
    if not time_col or not status_col:
        return ParsedOcm([], len(rows), len(rows), headers, time_col, status_col)
    for r in rows:
        ts = parse_ts(r.get(time_col))
        state = classify(r.get(status_col))
        if ts is None or state is None:
            skipped += 1
            continue
        ppm = None
        if ppm_col:
            try:
                ppm = float(str(r.get(ppm_col)).replace(",", "."))
            except (TypeError, ValueError):
                ppm = None
        records.append((ts, state, ppm, str(r.get(status_col))[:200]))
    return ParsedOcm(records, len(rows), skipped, headers, time_col, status_col)


def read_file(filename: str, content: bytes) -> list[dict]:
    """CSV (comma, semicolon or tab separated) or xlsx -> list of row dicts keyed by the header row."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            return []
        keys = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(header)]
        return [dict(zip(keys, row)) for row in it if any(c not in (None, "") for c in row)]
    text = content.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [dict(r) for r in csv.DictReader(io.StringIO(text), dialect=dialect)]


def operations(records: list) -> list[tuple]:
    """Discharge operations (start, end) from time-ordered (time, discharging, ...) records: a run of discharging
    records, ended by the first record that is not discharging (or by the last discharging record of the data)."""
    ops, start, last = [], None, None
    for rec in sorted(records, key=lambda r: r[0]):
        ts, on = rec[0], rec[1]
        if on:
            if start is None:
                start = ts
            last = ts
        elif start is not None:
            ops.append((start, ts))
            start = last = None
    if start is not None:
        ops.append((start, last))
    return ops


def overlaps(entry_start: datetime, entry_end: datetime, op: tuple, window_min: float) -> bool:
    """The entry's period, widened by the window each side, overlaps the recorded operation."""
    w = timedelta(minutes=window_min)
    return entry_start - w <= op[1] and op[0] <= entry_end + w
